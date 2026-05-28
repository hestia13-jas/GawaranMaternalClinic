import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from pptx import Presentation


SCRIPT_ROOT = Path(__file__).resolve().parents[1]


def add_kv_table(doc: Document, rows: list[tuple[str, str]]):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light List Accent 1"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Item"
    hdr_cells[1].text = "Details"
    for k, v in rows:
        r = table.add_row().cells
        r[0].text = k
        r[1].text = v
    doc.add_paragraph("")


def build_docx(out_path: Path):
    today = datetime.date.today().strftime("%B %d, %Y")
    project_title = "Gawaran Maternal Clinic"
    student_1 = "Bryan Longalong"
    student_2 = "Student Name 2"
    student_3 = "Student Name 3"

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # TITLE PAGE (as required by guideline PDF)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("TITLE PAGE\n\n")
    run.bold = True
    run.font.size = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{project_title}\n")
    r.bold = True
    r.font.size = Pt(26)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Prepared by:\n").bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{student_1}\n{student_2}\n{student_3}\n")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "Submitted to the Faculty of the Department of Computer Studies in Cavite State University –\n"
        "Imus, Cavite\n"
        "In partial Fulfilment of the Requirements for the Degree Bachelor of Science in Information\n"
        "Technology\n\n"
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(today)

    doc.add_page_break()

    # CHAPTER 1 — INTRODUCTION (sections 1–8)
    doc.add_heading("CHAPTER 1 — INTRODUCTION", level=1)

    doc.add_heading("1. Introduction", level=2)
    doc.add_paragraph(
        f"{project_title} is a full-stack maternity clinic web application built with Node.js (Express) and static "
        "HTML/CSS/JavaScript pages. It integrates with Supabase for authentication and data storage, enabling "
        "role-based clinic operations and patient access."
    )
    doc.add_paragraph("Why it is needed / current problems addressed:", style="List Bullet")
    for bullet in [
        "Clinic operations often rely on manual encoding and paper-based tracking, which increases delays and errors.",
        "Lack of centralized dashboards can reduce visibility into appointments, records, and clinic workload.",
        "Security risks increase without standardized access control, audit logs, and account protections.",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("2. Project Context", level=2)
    doc.add_paragraph("Existing workflow (typical):", style="List Bullet")
    for bullet in [
        "Patient registration and appointment scheduling handled manually.",
        "Staff coordination and notifications are done through ad-hoc messaging.",
        "Reports and audit trails are difficult to consolidate without a central system.",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")
    doc.add_paragraph("Current issues:", style="List Bullet")
    for bullet in ["Manual process", "Delays", "Human errors", "No real-time monitoring"]:
        doc.add_paragraph(bullet, style="List Bullet")
    doc.add_paragraph(
        "Automation is needed to provide a single portal that standardizes workflows, improves tracking, and "
        "adds security controls appropriate for a healthcare context."
    )

    doc.add_heading("3. Objectives of the Study", level=2)
    doc.add_paragraph("General Objective", style="List Bullet")
    doc.add_paragraph(
        "To develop a web-based maternity clinic management system that supports role-based access, "
        "centralized dashboards, and secure authentication.",
        style="List Bullet",
    )
    doc.add_paragraph("Specific Objectives", style="List Bullet")
    for bullet in [
        "Provide role-based portals for admin, doctor, nurse, staff, and patient users.",
        "Automate key transactions (registration, login, appointment visibility) and module access.",
        "Generate dashboard statistics and charts for monitoring clinic operations.",
        "Improve security through account lockout, 2FA (OTP), audit logs, and protective HTTP headers.",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("4. Purpose and Description", level=2)
    doc.add_paragraph("Features, users, benefits, and modules:", style="List Bullet")
    for bullet in [
        "Users: Admin, Doctor, Nurse, Staff, Patient.",
        "Benefits: improved visibility, reduced manual error, faster reporting, stronger security posture.",
        "Modules exposed by role: Patient Management, Medical Records, Laboratory & Diagnostics, Administration & Security, Communication & Notifications, Reporting & Analytics.",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")
    doc.add_paragraph("Key features (from the project):", style="List Bullet")
    for bullet in [
        "Landing page with service cards and compliance badges (PRC/HIPAA/RA 10173).",
        "Authentication: signup, login, forgot password, reset password, inline validation.",
        "2FA OTP via email (nodemailer SMTP) when enabled per user profile.",
        "Dashboards: admin/doctor/patient widgets and charts via `/api/dashboard/stats`.",
        "Audit logs endpoint for admins (`/api/audit`).",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("5. Time and Place of the Study", level=2)
    doc.add_paragraph("Timeline (high-level):", style="List Bullet")
    for bullet in [
        "Planning: requirements and feature outline",
        "Design: UI pages and API module design",
        "Development: server routes, Supabase schema, front-end pages",
        "Testing: local run, validation, role-based navigation checks",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")
    doc.add_paragraph("Place/Location:", style="List Bullet")
    doc.add_paragraph("Gawaran Maternal Clinic (deployment target) and development environment (local machine).", style="List Bullet")
    doc.add_paragraph("Respondents:", style="List Bullet")
    doc.add_paragraph("Clinic staff and patients (intended users).", style="List Bullet")

    doc.add_heading("6. Scope and Limitation", level=2)
    doc.add_paragraph("Scope (what the system CAN do):")
    for bullet in [
        "Serve a role-based portal for clinic stakeholders.",
        "Authenticate users, including optional OTP-based 2FA.",
        "Show dashboard widgets/charts and module availability based on role.",
        "Record and retrieve audit logs for administrative review.",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")
    doc.add_paragraph("Limitations (what the system CANNOT do):")
    for bullet in [
        "Does not include a mobile app (web-only).",
        "Real-time monitoring is limited (no live websocket streaming by default).",
        "Full clinical workflows (e.g., detailed prenatal tracking) are outside the current scope.",
        "Availability of cloud backup depends on configured Supabase project and environment.",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("7. Conceptual Framework", level=2)
    doc.add_paragraph("IPO (Input–Process–Output) Model")
    doc.add_paragraph("Input:", style="List Bullet")
    for bullet in ["User credentials and profile data", "Appointments/records/lab entries (via database)", "Role assignments and security flags (2FA/lockout)"]:
        doc.add_paragraph(bullet, style="List Bullet")
    doc.add_paragraph("Process:", style="List Bullet")
    for bullet in ["Authenticate and authorize users (RBAC)", "Load dashboard statistics and module list", "Apply security controls (rate limit, CSRF, lockout, OTP verification)", "Record audit logs"]:
        doc.add_paragraph(bullet, style="List Bullet")
    doc.add_paragraph("Output:", style="List Bullet")
    for bullet in ["Role-based portal pages", "Dashboard widgets/charts", "Audit log entries and security responses"]:
        doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("8. Definition of Terms", level=2)
    add_kv_table(
        doc,
        [
            ("Frontend", "Client-side pages (HTML/CSS/JS) served from `public/`."),
            ("Backend", "Node.js Express server (`server.js`) providing APIs and static hosting."),
            ("Supabase", "Managed backend (Auth + Postgres) used for users and data tables."),
            ("RBAC", "Role-based access control; permissions depend on user role."),
            ("RLS", "Row Level Security (Postgres policies) used to restrict table access."),
            ("OTP / 2FA", "One-Time Password used as a second authentication factor via email."),
            ("CSRF", "Cross-Site Request Forgery protection using server-issued tokens."),
        ],
    )

    doc.add_page_break()

    # CHAPTER 2 — REVIEW OF RELATED LITERATURE
    doc.add_heading("CHAPTER 2 — REVIEW OF RELATED LITERATURE", level=1)
    doc.add_heading("Related Literature", level=2)
    doc.add_paragraph(
        "This section summarizes general concepts supporting the project, such as clinic information systems, "
        "role-based access control, audit logging, and multi-factor authentication in web applications."
    )
    doc.add_heading("Related Studies", level=2)
    doc.add_paragraph(
        "This section discusses similar studies and systems used in healthcare management and how the proposed "
        "system aligns with best practices in security and usability."
    )

    doc.add_page_break()

    # CHAPTER 3 — METHODOLOGY
    doc.add_heading("CHAPTER 3 — METHODOLOGY", level=1)
    doc.add_heading("Development Model", level=2)
    doc.add_paragraph("Agile Prototyping", style="List Bullet")
    doc.add_paragraph(
        "The project follows an iterative approach: planning, design, development, and testing in short cycles, "
        "allowing rapid feedback on UI pages and API behavior."
    )
    doc.add_heading("Data Gathering", level=2)
    for bullet in ["Interview (clinic staff)", "Observation (existing clinic workflow)", "Survey (user feedback)"]:
        doc.add_paragraph(bullet, style="List Bullet")
    doc.add_heading("System Design Artifacts (high-level)", level=2)
    for bullet in ["Use Case Diagram (planned)", "ERD (based on Supabase schema)", "System Architecture (client/server/database)", "Process Flow (authentication + portal navigation)"]:
        doc.add_paragraph(bullet, style="List Bullet")

    doc.add_page_break()

    # CHAPTER 4 — SYSTEM DESIGN & IMPLEMENTATION
    doc.add_heading("CHAPTER 4 — SYSTEM DESIGN & IMPLEMENTATION", level=1)
    doc.add_heading("4.1 System Architecture", level=2)
    doc.add_paragraph("Frontend (Static Web Pages):", style="List Bullet")
    doc.add_paragraph("Pages under `public/` such as `index.html`, `login.html`, `portal.html`, and role dashboards.", style="List Bullet")
    doc.add_paragraph("Backend (Express API + Static Host):", style="List Bullet")
    doc.add_paragraph("`server.js` serves `/api/*` routes and hosts static assets.", style="List Bullet")
    doc.add_paragraph("Database (Supabase):", style="List Bullet")
    doc.add_paragraph("Postgres tables defined in `supabase/schema.sql` with Auth + RLS policies.", style="List Bullet")
    doc.add_paragraph("Email:", style="List Bullet")
    doc.add_paragraph("OTP emails sent via nodemailer SMTP; password reset handled via Supabase Auth.", style="List Bullet")

    doc.add_heading("4.2 Features of the System", level=2)
    for bullet in [
        "Authentication and profile management (register/login/reset).",
        "Role-based portal navigation and module availability.",
        "Dashboard statistics and charts by role.",
        "Security features: lockout, 2FA OTP, audit logs, CSRF checks, rate limiting.",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("4.3 Technologies Used", level=2)
    add_kv_table(
        doc,
        [
            ("Node.js", "Server runtime (>= 18)."),
            ("Express", "Web framework for API routes and static hosting."),
            ("Supabase", "Auth and Postgres database service."),
            ("HTML/CSS/JavaScript", "Frontend pages and interactivity."),
            ("Helmet", "Security headers and CSP."),
            ("express-rate-limit", "Rate limiting for auth endpoints."),
            ("nodemailer", "SMTP emails for OTP (2FA)."),
        ],
    )

    doc.add_heading("4.4 Screenshots and Description", level=2)
    doc.add_paragraph("Login Page", style="List Bullet")
    doc.add_paragraph("Screenshot: (capture during demo) — Provides login with optional OTP verification.", style="List Bullet")
    doc.add_paragraph("Portal Page", style="List Bullet")
    doc.add_paragraph("Screenshot: (capture during demo) — Central portal with role-based menus and modules.", style="List Bullet")
    doc.add_paragraph("Dashboard Widgets", style="List Bullet")
    doc.add_paragraph("Screenshot: (capture during demo) — Widgets and charts loaded from `/api/dashboard/stats`.", style="List Bullet")

    doc.add_page_break()

    # APPENDICES (as required by guideline PDF)
    doc.add_heading("APPENDICES", level=1)
    doc.add_heading("Appendix A — Gantt Chart", level=2)
    doc.add_paragraph("Gantt chart: (attach project timeline chart here).")
    doc.add_heading("Appendix B — Sample Source Code", level=2)
    doc.add_paragraph("Important modules:", style="List Bullet")
    for bullet in ["`server.js` (API + security middleware)", "`routes/auth.js` (auth + OTP)", "`supabase/schema.sql` (tables + RLS)"]:
        doc.add_paragraph(bullet, style="List Bullet")
    doc.add_heading("Appendix C — GitHub Repository", level=2)
    doc.add_paragraph("Repository link: (insert public repository link here).")
    doc.add_heading("Appendix D — Diagrams", level=2)
    doc.add_paragraph("Use Case Diagram / ERD / Architecture diagram: (attach diagrams here).")
    doc.add_heading("Appendix E — Evaluation Forms", level=2)
    doc.add_paragraph("Evaluation forms: (attach forms here).")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def add_bullets(slide, title: str, bullets: list[str]):
    layout = slide.slide_layouts[1]  # title + content
    s = slide.slides.add_slide(layout)
    s.shapes.title.text = title
    tf = s.shapes.placeholders[1].text_frame
    tf.clear()
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.level = 0
    return s


def build_pptx(out_path: Path):
    prs = Presentation()

    # Title slide
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = "Gawaran Maternal Clinic"
    subtitle = title_slide.placeholders[1]
    subtitle.text = (
        "Group Members: Bryan Longalong, Student 2, Student 3\n"
        "Subject: (Insert Subject)\n"
        "Adviser: (Insert Adviser)\n"
        "Cavite State University – Imus, Cavite\n"
        f"{datetime.date.today().strftime('%B %d, %Y')}"
    )

    # Slide 2 — Introduction (per guideline PDF)
    add_bullets(
        prs,
        "Introduction",
        [
            "A full-stack maternity clinic web application for clinic operations and patient access",
            "Purpose: centralize workflows, dashboards, and role-based portal access",
            "Built with Node.js (Express) + static web pages + Supabase backend services",
        ],
    )

    # Slide 3 — Current Problems
    add_bullets(
        prs,
        "Current Problems",
        [
            "Manual process",
            "Delays",
            "Human errors",
            "Long waiting time",
            "No real-time monitoring",
        ],
    )

    # Slide 4 — Objectives of the Study
    add_bullets(
        prs,
        "Objectives of the Study",
        [
            "General Objective: Develop a secure web-based clinic management portal",
            "Specific Objectives:",
            "• Automate key transactions (registration, authentication, portal access)",
            "• Generate dashboards and basic reports for monitoring",
            "• Improve monitoring through centralized widgets and charts",
            "• Reduce manual errors via validation and standardized workflows",
        ],
    )

    # Slide 5 — Scope and Limitations
    add_bullets(
        prs,
        "Scope and Limitations",
        [
            "Scope (CAN do): role-based portals (admin/doctor/nurse/staff/patient), dashboards, auth + 2FA, audit logs",
            "Limitations (CANNOT do): no mobile app, limited real-time updates, advanced clinical workflows not included",
        ],
    )

    # Slide 6 — System Architecture (helpful for defense; keeps minimal text)
    add_bullets(
        prs,
        "System Architecture",
        [
            "Frontend: `public/` HTML/CSS/JS pages (portal + dashboards)",
            "Backend: `server.js` Express API under `/api/*`",
            "Database/Auth: Supabase (Postgres + Auth + RLS policies)",
            "Security: Helmet+CSP, rate limits, CSRF token checks, OTP (2FA), audit logs",
        ],
    )

    # Slide 7 — Methodology
    add_bullets(
        prs,
        "Methodology",
        [
            "Agile Prototyping",
            "Planning → Design → Development → Testing",
            "Iterative improvements for UI validation, role-based modules, and security controls",
        ],
    )

    # Slide 8 — System Demo
    add_bullets(
        prs,
        "System Demo",
        [
            "Live demo flow:",
            "• Open landing page",
            "• Sign up / login",
            "• (Optional) OTP verification for 2FA-enabled account",
            "• Open portal and view role-based modules + dashboard widgets",
        ],
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)


def main():
    # Write outputs to the *current working directory*'s docs folder so it matches
    # wherever the user runs the project from (Desktop vs Projects path).
    out_dir = Path.cwd() / "docs"
    out_dir.mkdir(parents=True, exist_ok=True)

    # Use new filenames to avoid PermissionError when older files are open/locked.
    docx_path = out_dir / "GawaranMaternalClinic_Documentation_Guidelines.docx"
    pptx_path = out_dir / "GawaranMaternalClinic_Presentation_Guidelines.pptx"
    build_docx(docx_path)
    build_pptx(pptx_path)
    print(f"Wrote {docx_path}")
    print(f"Wrote {pptx_path}")


if __name__ == "__main__":
    main()

